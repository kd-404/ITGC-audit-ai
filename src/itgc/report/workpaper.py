"""Builds the audit workpaper: an Excel file with one row per finding
(condition/criteria/cause/effect/recommendation/evidence - the standard ITGC
finding structure), plus supporting sheets an auditor would expect to see:
what evidence was ingested and how, and what open enquiries exist.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from ..models import Finding, Enquiry

SEVERITY_FILL = {
    "Critical": "C00000",
    "High": "E26B0A",
    "Medium": "FFC000",
    "Low": "FFFF99",
}
HEADER_FILL = "1F4E78"


def _style_header(ws, row=1):
    for cell in ws[row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(vertical="top", wrap_text=True)


def _autosize(ws, widths: dict[int, int]):
    for col_idx, width in widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def build_workpaper(out_path: Path, findings: list[Finding], enquiries: list[Enquiry],
                     file_classifications: list[dict], engagement_name: str, as_of: str) -> Path:
    wb = openpyxl.Workbook()

    # --- Findings ---
    ws = wb.active
    ws.title = "Findings"
    headers = ["ID", "Severity", "Type", "System", "Subject", "Condition", "Criteria",
               "Cause", "Effect", "Recommendation", "Evidence", "Confidence", "Narrative"]
    ws.append(headers)
    _style_header(ws)
    for f in findings:
        ws.append([
            f.finding_id, f.severity.value, f.finding_type.value, f.system,
            f"{f.subject_person or ''} ({f.subject_email or 'n/a'})",
            f.condition, f.criteria, f.cause, f.effect, f.recommendation,
            "; ".join(f.evidence), f.confidence, f.narrative or "",
        ])
        row = ws.max_row
        fill = SEVERITY_FILL.get(f.severity.value)
        if fill:
            ws.cell(row=row, column=2).fill = PatternFill("solid", fgColor=fill)
        for c in range(1, len(headers) + 1):
            ws.cell(row=row, column=c).alignment = Alignment(vertical="top", wrap_text=True)
    _autosize(ws, {1: 8, 2: 10, 3: 30, 4: 10, 5: 26, 6: 45, 7: 35, 8: 30, 9: 32, 10: 32, 11: 35, 12: 10, 13: 45})
    ws.freeze_panes = "A2"

    # --- Enquiries ---
    ws2 = wb.create_sheet("Enquiries")
    ws2.append(["ID", "Blocking", "Topic", "Question", "Context", "Answer"])
    _style_header(ws2)
    for e in enquiries:
        ws2.append([e.enquiry_id, "Yes" if e.blocking else "No", e.topic, e.question, e.context, e.answer or ""])
        for c in range(1, 7):
            ws2.cell(row=ws2.max_row, column=c).alignment = Alignment(vertical="top", wrap_text=True)
    _autosize(ws2, {1: 8, 2: 9, 3: 26, 4: 55, 5: 30, 6: 30})
    ws2.freeze_panes = "A2"

    # --- Evidence / ingestion audit trail ---
    ws3 = wb.create_sheet("Evidence Index")
    ws3.append(["File", "Classified As", "Confidence"])
    _style_header(ws3)
    for c in file_classifications:
        ws3.append([c["file"], c["evidence_type"], c["confidence"]])
    _autosize(ws3, {1: 45, 2: 20, 3: 12})
    ws3.freeze_panes = "A2"

    # --- Cover / summary ---
    ws4 = wb.create_sheet("Summary", 0)
    ws4["A1"] = engagement_name
    ws4["A1"].font = Font(bold=True, size=14)
    ws4["A2"] = f"As of: {as_of}"
    ws4["A3"] = f"Total findings: {len(findings)}"
    counts = {}
    for f in findings:
        counts[f.severity.value] = counts.get(f.severity.value, 0) + 1
    r = 5
    ws4[f"A{r}"] = "Severity"
    ws4[f"B{r}"] = "Count"
    _style_header(ws4, row=r)
    for sev in ["Critical", "High", "Medium", "Low"]:
        r += 1
        ws4[f"A{r}"] = sev
        ws4[f"B{r}"] = counts.get(sev, 0)
    _autosize(ws4, {1: 24, 2: 12})

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path
